"""Feedback round 2026-10-22 (b), applied to v2.html. Rerunnable.

1. Vision and Foundation page: JSW Group and JSW Motors become two clearly separate panels
   (text untouched; only wrapper elements and CSS are added)
2. Chapter 5 annexures relinked to 'Ch 5 Annexures.xlsx' (5A-5D; the old 5E is gone), rendered with the same
   sheet renderer as annexures.py; Use / Used-by lines come from that workbook's Index sheet
3. 9B Certificate Template redesigned as a certificate; every word is the 9B sheet's own text

usage: python3 portal-v7/feedback_oct22b.py [Ch5-xlsx] [JSW_HR_Manual_Annexures-xlsx]
"""
import html, json, math, re, sys, warnings
from pathlib import Path
import openpyxl

warnings.filterwarnings('ignore')
ROOT = Path(__file__).resolve().parent.parent
V2 = ROOT / 'v2.html'
CH5 = Path(sys.argv[1]) if len(sys.argv) > 1 else Path.home() / 'Downloads/Ch 5 Annexures.xlsx'
ALL = Path(sys.argv[2]) if len(sys.argv) > 2 else Path.home() / 'Downloads/JSW_HR_Manual_Annexures.xlsx'
esc = lambda s: html.escape(s, quote=False)

h = V2.read_text(encoding='utf-8')

# ---------------------------------------------------------------- 1. two parts
a = h.index('<section class="chap" id="fm-i">')
b = h.index('</section>', a)
sec = h[a:b]
if 'vf-part' not in sec:
    g0 = '<h3 class="sub-title subhead" style="margin-top:0">JSW Group</h3>'
    m0 = '<h3 class="sub-title subhead">JSW Motors</h3>'
    assert g0 in sec and m0 in sec
    gi, mi = sec.index(g0), sec.index(m0)
    end = sec.rindex('</div>')
    group, motors = sec[gi + len(g0):mi], sec[mi + len(m0):end]
    head = lambda t: f'<header class="vf-head"><h3 class="sub-title subhead">{t}</h3></header>'
    new = (f'<div class="vf-part vf-group">{head("JSW Group")}<div class="vf-body">{group}</div></div>'
           f'<div class="vf-part vf-motors">{head("JSW Motors")}<div class="vf-body">{motors}</div></div>')
    sec = sec[:gi] + new + sec[end:]
    h = h[:a] + sec + h[b:]

# ---------------------------------------------------------------- 2. chapter 5
TITLES, DESC = {}, {}
for r in openpyxl.load_workbook(CH5, data_only=True)['Index'].iter_rows(min_row=5, values_only=True):
    if r[1] and r[2]:
        code = str(r[1]).strip()
        TITLES[code] = str(r[2]).strip()
        DESC[code] = re.sub(r'\s+', ' ', str(r[3])).strip()
assert list(TITLES) == ['5A', '5B', '5C', '5D'], TITLES
SHEET = {c: next(s for s in openpyxl.load_workbook(CH5, read_only=True).sheetnames if s.startswith(c)) for c in TITLES}

src = (ROOT / 'portal-v7/annexures.py').read_text()
chunk = src[src.index('def fmt('):src.index('# 3D is a picture')]
G = {'openpyxl': openpyxl, 're': re, 'datetime': __import__('datetime'), 'html': html,
     'get_column_letter': openpyxl.utils.get_column_letter}
G['esc'] = lambda s: html.escape(s, quote=True)
G['wb'] = openpyxl.Workbook()
exec(chunk, G)


def render_sheet(code):
    wb = openpyxl.load_workbook(CH5)
    wv = openpyxl.load_workbook(CH5, data_only=True)
    G['wb'] = wb
    G['THEME'] = G['_theme_colors']()
    return G['render'](wb[SHEET[code]], wv[SHEET[code]])


m = re.search(r'<script id="jx-annex-docs">window\.ANNEX_DOCS = (\{.*?\});</script>', h, re.S)
docs = json.loads(m.group(1).replace('<\\/', '</'))
docs.pop('5E', None)
for code, t in TITLES.items():
    docs[code] = {'code': code, 'title': t, 'filename': f'{code}_{t.replace(" ", "_")}.doc', 'content': render_sheet(code)}

# ---------------------------------------------------------------- 3. certificate (9B)
CERT_TEXT = {c.coordinate: c.value for c in openpyxl.load_workbook(ALL)['9B_Certificate_Template']['B'] if c.value}
T = lambda k: esc(CERT_TEXT[k])
sig = [re.match(r'_+\s+(.*)', CERT_TEXT[k]).group(1) for k in ('B15', 'B16', 'B17')]
cat = CERT_TEXT['B8'].split(':', 1)
cit = CERT_TEXT['B9'].split(':', 1)
dte = CERT_TEXT['B12'].split(':', 1)
assert CERT_TEXT['B2'] == 'CERTIFICATE OF EXCELLENCE'


def scallop(cx, cy, r1, r2, n):
    pts = []
    for i in range(n * 2):
        r = r1 if i % 2 == 0 else r2
        t = math.pi * i / n
        pts.append(f'{cx + r * math.sin(t):.1f},{cy - r * math.cos(t):.1f}')
    return ' '.join(pts)


GOLD = '#a8863a'
corner = (f'<g fill="none" stroke="{GOLD}" stroke-linecap="round" stroke-linejoin="round">'
          '<path d="M36 36 C36 92 70 126 126 126" stroke-width="2.2"/><path d="M36 36 C92 36 126 70 126 126" stroke-width="2.2"/>'
          '<path d="M48 48 C48 84 68 104 104 104" stroke-width="1.4"/><path d="M48 48 C84 48 104 68 104 104" stroke-width="1.4"/>'
          '<path d="M60 60 C60 76 68 84 84 84 C94 84 98 76 94 70 C90 66 84 70 86 74" stroke-width="1.6"/>'
          '<path d="M60 60 C76 60 84 68 84 84" stroke-width="1.2"/>'
          '<path d="M126 126 C136 126 146 120 150 110 M126 126 C126 136 120 146 110 150" stroke-width="1.6"/>'
          f'<circle cx="36" cy="36" r="4" fill="{GOLD}" stroke="none"/><circle cx="154" cy="104" r="2.500" fill="{GOLD}" stroke="none"/>'
          f'<circle cx="104" cy="154" r="2.500" fill="{GOLD}" stroke="none"/></g>')
edge = (f'<g fill="none" stroke="{GOLD}" stroke-width="1.600" stroke-linecap="round">'
        '<path d="M-26 0 C-18 -10 -8 -10 0 0 C8 10 18 10 26 0"/><path d="M-26 0 C-34 4 -40 -2 -34 -6"/><path d="M26 0 C34 4 40 -2 34 -6"/>'
        f'<circle r="3" fill="{GOLD}" stroke="none"/></g>')
cs = f'<g transform="translate(8 8) scale(.6)">{corner}</g>'
frame = (f'<svg class="cert-art" viewBox="0 0 600 848" preserveAspectRatio="none" aria-hidden="true" focusable="false">'
         f'<rect x="22" y="22" width="556" height="804" fill="none" stroke="{GOLD}" stroke-width="2.400"/>'
         f'<rect x="31" y="31" width="538" height="786" fill="none" stroke="{GOLD}" stroke-width="0.900"/>'
         f'<g>{cs}</g><g transform="translate(600 0) scale(-1 1)">{cs}</g>'
         f'<g transform="translate(0 848) scale(1 -1)">{cs}</g><g transform="translate(600 848) scale(-1 -1)">{cs}</g>'
         f'<g transform="translate(300 26)">{edge}</g><g transform="translate(300 822) scale(1 -1)">{edge}</g>'
         f'<g transform="translate(26 424) rotate(-90)">{edge}</g><g transform="translate(574 424) rotate(90)">{edge}</g></svg>')
seal = ('<svg class="cert-seal" viewBox="0 0 120 150" aria-hidden="true" focusable="false"><defs>'
        '<radialGradient id="cs-g" cx="40%" cy="35%" r="75%"><stop offset="0" stop-color="#f7e7a6"/><stop offset=".55" stop-color="#c9a44a"/><stop offset="1" stop-color="#8f6d24"/></radialGradient>'
        '<linearGradient id="cs-r" x1="0" y1="0" x2="1" y2="1"><stop offset="0" stop-color="#e9cf7e"/><stop offset="1" stop-color="#a98332"/></linearGradient></defs>'
        '<path d="M34 92 22 146 46 134 60 148 66 100z" fill="url(#cs-r)"/><path d="M86 92 98 146 74 134 60 148 54 100z" fill="url(#cs-r)" opacity=".92"/>'
        f'<polygon points="{scallop(60, 58, 54, 48, 18)}" fill="url(#cs-g)"/>'
        '<circle cx="60" cy="58" r="40" fill="none" stroke="#f8eab0" stroke-width="1.500"/><circle cx="60" cy="58" r="33" fill="url(#cs-g)" stroke="#8f6d24" stroke-width=".8"/>'
        '<path d="M60 36 64.500 52 81 52 67.500 61.500 72.500 77 60 67.500 47.500 77 52.500 61.500 39 52 55.500 52z" fill="#f8eab0" opacity=".85"/></svg>')

CERT = (
    '<div class="cert-wrap"><div class="cert">' + frame +
    '<div class="cert-in">'
    f'<div class="cert-org">{T("B19")}</div>'
    f'<h2 class="cert-title"><span>CERTIFICATE</span><em>OF EXCELLENCE</em></h2>'
    f'<div class="cert-lbl">{T("B3")}</div>'
    f'<div class="cert-name">{T("B4")}</div>'
    f'<p class="cert-p">{T("B5")}</p>'
    f'<p class="cert-p">{T("B7")}</p>'
    f'<div class="cert-award"><span class="cert-k">{esc(cat[0])}:</span> <strong>{esc(cat[1].strip())}</strong></div>'
    f'<p class="cert-cite"><span class="cert-k">{esc(cit[0])}:</span> {esc(cit[1].strip())}</p>'
    f'<div class="cert-date"><span class="cert-k">{esc(dte[0])}:</span> {esc(dte[1].strip())}</div>'
    + seal +
    f'<div class="cert-sigs-lbl">{T("B14")}</div>'
    '<div class="cert-sigs">' + ''.join(f'<div class="cert-sig"><i></i><span>{esc(x)}</span></div>' for x in sig) + '</div>'
    f'<p class="cert-quote">{T("B20")}</p>'
    '</div></div></div>')
docs['9B'] = {'code': '9B', 'title': 'Certificate Template', 'filename': '9B_Certificate_Template.doc', 'content': CERT}

docs = {k: docs[k] for k in sorted(docs, key=lambda c: (int(re.match(r'\d+', c).group()), c))}
new_js = '<script id="jx-annex-docs">window.ANNEX_DOCS = ' + json.dumps(docs, ensure_ascii=False).replace('</', '<\\/') + ';</script>'
m = re.search(r'<script id="jx-annex-docs">window\.ANNEX_DOCS = (\{.*?\});</script>', h, re.S)
h = h[:m.start()] + new_js + h[m.end():]

# ---- rows
EYE = ('<svg width="13" height="13" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2.2">'
       '<path d="M1 12s4-8 11-8 11 8 11 8-4 8-11 8-11-8-11-8z"></path><circle cx="12" cy="12" r="3"></circle></svg>')


def desc_html(text):
    out = []
    for p in (p.strip() for p in text.split('|')):
        k, _, v = p.partition(':')
        out.append(f'<span><b>{esc(k)}:</b> {esc(v.strip())}</span>')
    return '<div class="annex-file-desc">' + ''.join(out) + '</div>'


def item(code, appendix):
    t = f'{code}. {TITLES[code]}'
    left = (f'<div class="annex-file-left"><div class="annex-file-icon">📄</div><div class="annex-file-text"><div class="annex-file-name">{t}</div>{desc_html(DESC[code])}</div></div>'
            if appendix else f'<div class="annex-file-left"><div class="annex-file-icon">📄</div><div class="annex-file-name">{t}</div></div>')
    return (f'<div{f" id=\"annex-{code}\"" if appendix else ""} class="annex-file-item" role="button" tabindex="0" title="View {t}" '
            f'onclick="openAnnexViewer(event, \'{code}\')" onkeydown="if(event.key===\'Enter\'||event.key===\' \'){{openAnnexViewer(event, \'{code}\')}}">'
            f'{left}<div class="annex-view-badge">{EYE} View File</div></div>')


BLOCK = re.compile(r'<!--annex-list--><div class="annex-file-list">(?:(?!<!--/annex-list-->).)*?openAnnexViewer\(event, \'5A\'\)(?:(?!<!--/annex-list-->).)*?</div><!--/annex-list-->', re.S)
blocks = list(BLOCK.finditer(h))
assert len(blocks) == 2, len(blocks)
for mm in reversed(blocks):  # appendix is the later one
    appendix = 'id="annex-5A"' in mm.group(0)
    new = '<!--annex-list--><div class="annex-file-list">' + ''.join(item(c, appendix) for c in TITLES) + '</div><!--/annex-list-->'
    h = h[:mm.start()] + new + h[mm.end():]

# ---------------------------------------------------------------- css
CSS = r'''<style id="jx-oct22b-css">
/* Vision and Foundation: two separate panels, one shared grey */
.jx-chapter #fm-i .vf-part{border-radius:22px;overflow:hidden;margin:0 0 36px;border:1.5px solid #8d8d90;background:#d4d4d2;color:#1b1b1d;box-shadow:0 8px 24px rgba(0,0,0,.10)}
.jx-chapter #fm-i .vf-head{padding:26px 34px 22px;display:flex;align-items:center;gap:16px;background:rgba(0,0,0,.07);border-bottom:1px solid #8d8d90}
.jx-chapter #fm-i .vf-head::before{content:'';flex:none;width:6px;align-self:stretch;border-radius:3px;background:#1b1b1d}
.jx-chapter #fm-i .vf-head h3.sub-title{margin:0 !important;padding:0 !important;border:0 !important;font-size:clamp(26px,3.4vw,38px) !important;line-height:1.1 !important;letter-spacing:-.025em !important;font-weight:800 !important;color:#111 !important}
.jx-chapter #fm-i .vf-body{padding:8px 34px 30px}
.jx-chapter #fm-i .vf-body>*:first-child{margin-top:20px}
.jx-chapter #fm-i .vf-body h4{margin-top:30px}
.jx-chapter #fm-i .vf-body :is(h4,h5){color:#111 !important}
.jx-chapter #fm-i .vf-body p{color:#2a2a2d !important}
.jx-chapter #fm-i .vf-body :is(h4,h5,p,hr){border-color:#9a9a9d !important}
@media(max-width:620px){.jx-chapter #fm-i .vf-head{padding:20px 20px 16px}.jx-chapter #fm-i .vf-body{padding:4px 20px 22px}.jx-chapter #fm-i .vf-part{border-radius:16px}}

/* 9B certificate */
.cert-wrap{padding:6px 0 10px}
.cert{--g:#a8863a;position:relative;container-type:inline-size;max-width:640px;margin:0 auto;aspect-ratio:600/848;background:#f6f3ec;color:#262626;border-radius:4px;box-shadow:0 10px 34px rgba(60,45,10,.2);font-family:"Cormorant Garamond","Playfair Display",Georgia,"Times New Roman",serif}
.cert-art{position:absolute;inset:0;width:100%;height:100%}
.cert-in{position:relative;padding:12cqw 13cqw 9.5cqw;display:flex;flex-direction:column;align-items:center;text-align:center}
.cert-org{font-size:2.7cqw;line-height:1.5;letter-spacing:.14em;text-transform:uppercase;color:var(--g);font-weight:600;margin-top:0}
.cert-title{margin:4cqw 0 0;display:flex;flex-direction:column;align-items:center;font-weight:500}
.cert-title span{font-size:9.5cqw;line-height:1;letter-spacing:.02em;color:#232323}
.cert-title em{font-style:normal;font-size:3.3cqw;letter-spacing:.5em;margin-right:-.5em;color:var(--g);margin-top:1.4cqw}
.cert-lbl{margin-top:4.5cqw;font-size:2.9cqw;letter-spacing:.2em;text-transform:uppercase}
.cert-name{margin-top:2.5cqw;width:80%;padding:0 0 2cqw;border-bottom:1.5px solid var(--g);font-family:"Snell Roundhand","Apple Chancery","Lucida Handwriting","Brush Script MT","Segoe Script",cursive;font-size:7.2cqw;line-height:1.15;color:#1d1d1d}
.cert-p{margin:2.2cqw 0 0;font-size:2.85cqw;line-height:1.5;letter-spacing:.03em;max-width:84%;color:#333}
.cert-award{margin-top:3cqw;font-size:3.6cqw;letter-spacing:.03em}
.cert-award strong{font-weight:700;color:#232323}
.cert-k{color:var(--g);font-weight:600}
.cert-cite{margin:1.4cqw 0 0;font-size:2.7cqw;line-height:1.45;max-width:84%;font-style:italic;color:#444}
.cert-date{margin-top:2.2cqw;font-size:2.8cqw;letter-spacing:.04em}
.cert-sigs-lbl{margin-top:1.5cqw;font-size:2.3cqw;letter-spacing:.3em;text-transform:uppercase;color:var(--g);font-weight:600}
.cert-sigs{margin-top:2cqw;width:100%;display:grid;grid-template-columns:repeat(3,1fr);align-items:end;gap:3cqw}
.cert-sigs-one{grid-template-columns:1fr;justify-items:center;margin-top:1.8cqw}
.cert-sigs-one .cert-sig{width:44%}
.cert-sig{display:flex;flex-direction:column;align-items:center;gap:1cqw}
.cert-sig i{display:block;width:100%;height:0;border-top:1.5px solid var(--g);margin-top:4cqw}
.cert-sig{min-width:0}.cert-sig span{font-size:1.65cqw;letter-spacing:.03em;line-height:1.35;overflow-wrap:break-word;text-transform:uppercase;color:#333;font-weight:600}
.cert-seal{width:11cqw;height:auto;margin-top:2cqw}
.cert-quote{margin:2.4cqw 0 0;font-size:2.4cqw;line-height:1.45;font-style:italic;color:#555;max-width:80%}
</style>
'''
if 'id="jx-oct22b-css"' in h:
    h = re.sub(r'<style id="jx-oct22b-css">.*?</style>\n', lambda _: CSS, h, flags=re.S)
else:
    h = h.replace('<style id="jx-annex-docs-css">', CSS + '<style id="jx-annex-docs-css">', 1)

V2.write_text(h, encoding='utf-8')
print('ok; docs:', len(docs), '5x:', [k for k in docs if k.startswith('5')])
